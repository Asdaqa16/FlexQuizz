import random
import asyncio
import io
import logging
import os
from functools import wraps

from dotenv import load_dotenv
from google import genai
from pydantic import TypeAdapter
from pypdf import PdfReader

from AI_Layer.exceptions import (
    ConceptExtractionError,
    PDFProcessingError,
    QuestionGenerationError,
    ValidationFailedError,
)

from AI_Layer.models import (
    AnswerVerification,
    CalibrationEstimate,
    ChatTurn,
    Chunk,
    Citation,
    Concept,
    ConceptDraft,
    Difficulty,
    GroundedChatResponse,
    QuestionDraft,
    QuizReportData,
    ValidationResult,
)


logger = logging.getLogger(__name__)


# ============================================================
# ENVIRONMENT / GEMINI
# ============================================================

load_dotenv()

api_key = os.getenv("GEMINI_API_KEY")

try:
    if not api_key:
        raise ValueError("GEMINI_API_KEY is not set")

    client = genai.Client(
        api_key=api_key
    )

except ValueError as e:
    logger.warning(
        "Failed to initialize Gemini Client: %s",
        e,
    )

    client = None


# ============================================================
# CONSTANTS
# ============================================================

MODEL_TEXT = "gemini-3.1-flash-lite"

MAX_QUIZ_CONCEPTS = 8


# ============================================================
# RETRY HELPER
# ============================================================

def with_retry(
    retries: int = 0,
    backoff: float = 1.5,
    custom_exc=None,
):
    def decorator(func):

        @wraps(func)
        async def wrapper(*args, **kwargs):

            for attempt in range(retries + 1):

                try:
                    return await func(
                        *args,
                        **kwargs,
                    )

                except Exception as e:

                    if attempt == retries:

                        if custom_exc:
                            raise custom_exc(
                                f"Failed after {retries} retries "
                                f"in {func.__name__}: {e}"
                            ) from e

                        raise

                    await asyncio.sleep(
                        backoff ** attempt
                    )

        return wrapper

    return decorator


# ============================================================
# PDF EXTRACTION
# ============================================================

@with_retry(retries=1, custom_exc=PDFProcessingError)
async def extract_and_chunk(
    pdf_bytes: bytes,
    source_filename: str = "",
) -> list[Chunk]:

    try:

        reader = PdfReader(
            io.BytesIO(pdf_bytes)
        )

        chunks = []

        current_words = []

        chunk_index = 0

        start_page = 1

        TARGET_WORDS = 1200

        OVERLAP_WORDS = 150

        for page_num, page in enumerate(
            reader.pages,
            1,
        ):

            text = page.extract_text()

            if not text:
                continue

            words = text.split()

            for word in words:

                current_words.append(
                    word
                )

                if len(current_words) >= TARGET_WORDS:

                    page_range = f"{start_page}-{page_num}"

                    chunks.append(
                        Chunk(
                            chunk_index=chunk_index,
                            text=" ".join(
                                current_words
                            ),
                            approx_page_range=page_range,
                            source_type="pdf",
                            locator=f"Page {page_range}",
                            source_filename=source_filename,
                        )
                    )

                    chunk_index += 1

                    current_words = (
                        current_words[
                            -OVERLAP_WORDS:
                        ]
                    )

                    start_page = page_num

        if current_words:

            page_range = f"{start_page}-{len(reader.pages)}"

            chunks.append(
                Chunk(
                    chunk_index=chunk_index,
                    text=" ".join(
                        current_words
                    ),
                    approx_page_range=page_range,
                    source_type="pdf",
                    locator=f"Page {page_range}",
                    source_filename=source_filename,
                )
            )

        # --------------------------------------------------------
        # OCR Fallback for scanned / handwritten PDFs
        # --------------------------------------------------------
        
        total_words = sum(len(chunk.text.split()) for chunk in chunks)
        
        if total_words < 20:
            
            if client is None:
                raise PDFProcessingError(
                    "Gemini client is not initialized for OCR fallback."
                )
            
            print(f"[DEBUG] total_words={total_words}. Falling back to Gemini OCR.")
            
            prompt = (
                "Please transcribe all the readable text from this document "
                "as accurately as possible. Output only the transcribed text."
            )
            
            try:
                response = await client.aio.models.generate_content(
                    model=MODEL_TEXT,
                    contents=[
                        prompt,
                        genai.types.Part.from_bytes(
                            data=pdf_bytes,
                            mime_type="application/pdf",
                        )
                    ]
                )
                print(f"[DEBUG] Gemini OCR response text: {repr(response.text)}")
                transcribed_text = response.text or ""
            except Exception as gemini_e:
                print(f"[DEBUG] Gemini OCR failed: {repr(gemini_e)}")
                transcribed_text = ""

            return extract_text_and_chunk(transcribed_text, source_filename)

        return chunks

    except Exception as e:

        raise PDFProcessingError(
            f"Failed to process PDF: {e}"
        ) from e


# ============================================================
# TEXT EXTRACTION
# ============================================================

def extract_text_and_chunk(
    text: str,
    source_filename: str = "",
) -> list[Chunk]:

    text = text.strip()

    if not text:
        return []

    return [
        Chunk(
            chunk_index=0,
            text=text,
            approx_page_range="1",
            source_type="text",
            locator="Document",
            source_filename=source_filename,
        )
    ]


# ============================================================
# SLIDE DECK EXTRACTION (.pptx)
# ============================================================

def extract_slide_deck(
    pptx_bytes: bytes,
    source_filename: str = "",
) -> list[Chunk]:

    from pptx import Presentation

    prs = Presentation(io.BytesIO(pptx_bytes))

    chunks = []

    for slide_number, slide in enumerate(prs.slides, 1):

        texts = []

        for shape in slide.shapes:

            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())

            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(
                        " | ".join(
                            cell.text.strip()
                            for cell in row.cells
                        )
                    )

        slide_text = "\n".join(texts).strip()

        if not slide_text:
            continue

        chunks.append(
            Chunk(
                chunk_index=slide_number - 1,
                text=slide_text,
                approx_page_range=str(slide_number),
                source_type="slide",
                locator=f"Slide {slide_number}",
                source_filename=source_filename,
            )
        )

    return chunks


# ============================================================
# VIDEO EXTRACTION (native Gemini multimodal ingestion)
# ============================================================

@with_retry(retries=1, custom_exc=PDFProcessingError)
async def extract_video(
    video_bytes: bytes,
    mime_type: str = "video/mp4",
    source_filename: str = "",
) -> list[Chunk]:

    if client is None:
        raise PDFProcessingError(
            "Gemini client is not initialized for video ingestion."
        )

    prompt = (
        "Transcribe this lecture video into timestamped segments. "
        "Break it into natural topic segments of roughly 60-180 seconds "
        "each. For each segment, include what was said AND a description "
        "of any on-screen diagrams, slides, or written material. "
        "Respond as a JSON list of objects with fields "
        "'start_time' (format MM:SS), 'end_time' (format MM:SS), "
        "and 'text' (the transcribed/described content)."
    )

    response = await client.aio.models.generate_content(
        model=MODEL_TEXT,
        contents=[
            prompt,
            genai.types.Part.from_bytes(
                data=video_bytes,
                mime_type=mime_type,
            ),
        ],
        config=genai.types.GenerateContentConfig(
            response_mime_type="application/json",
        ),
    )

    import json

    try:
        segments = json.loads(response.text or "[]")
    except json.JSONDecodeError:
        segments = []

    chunks = []

    for index, segment in enumerate(segments):

        text = (segment.get("text") or "").strip()

        if not text:
            continue

        start = segment.get("start_time", "00:00")
        end = segment.get("end_time", start)

        chunks.append(
            Chunk(
                chunk_index=index,
                text=text,
                approx_page_range=f"{start}-{end}",
                source_type="video",
                locator=f"{start}-{end}",
                source_filename=source_filename,
            )
        )

    return chunks


# ============================================================
# CONCEPT EXTRACTION
# ============================================================

@with_retry(
    retries=1,
    custom_exc=ConceptExtractionError,
)
async def extract_concepts(
    chunks: list[Chunk],
) -> list[ConceptDraft]:

    if not chunks:
        return []

    if client is None:

        raise ConceptExtractionError(
            "Gemini client is not initialized. "
            "Check GEMINI_API_KEY."
        )

    MAX_SOURCE_WORDS = 9000

    source_words = []

    for chunk in chunks:

        source_words.extend(
            chunk.text.split()
        )

        if len(source_words) >= MAX_SOURCE_WORDS:
            break

    source_text = " ".join(
        source_words[:MAX_SOURCE_WORDS]
    )

    prompt = (
        "Extract the most important distinct educational "
        "concepts or topics from the following study material.\n\n"

        f"Return at most {MAX_QUIZ_CONCEPTS} concepts.\n"

        "Each concept must be meaningfully distinct.\n"

        "For each concept provide:\n"
        "- a short label\n"
        "- a one-line description\n"
        "- a parent 'topic': a broader subject heading that groups "
        "related concepts together (e.g. multiple concepts can share "
        "the same topic)\n\n"

        "Merge overlapping concepts instead of repeating them.\n\n"

        f"Source material:\n{source_text}"
    )

    response = await client.aio.models.generate_content(

        model=MODEL_TEXT,

        contents=prompt,

        config=genai.types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=list[ConceptDraft],
            temperature=0.1,
        ),
    )

    drafts = TypeAdapter(
        list[ConceptDraft]
    ).validate_json(
        response.text
    )

    drafts = drafts[
        :MAX_QUIZ_CONCEPTS
    ]

    for draft in drafts:
        draft.source_chunk_index = 0

    return drafts


# ============================================================
# CONCEPT DEDUPLICATION
# ============================================================

async def deduplicate_concepts(
    concept_drafts: list[ConceptDraft],
    chunks: list[Chunk] | None = None,
) -> list[Concept]:

    if not concept_drafts:
        return []

    chunks = chunks or []

    locator_by_index = {
        chunk.chunk_index: chunk.locator
        for chunk in chunks
    }

    final_concepts = []

    seen_labels = set()

    for draft in concept_drafts[
        :MAX_QUIZ_CONCEPTS
    ]:

        label = draft.label.strip()

        if not label:
            continue

        key = label.lower()

        if key in seen_labels:
            continue

        seen_labels.add(key)

        locator = locator_by_index.get(
            draft.source_chunk_index,
            "",
        )

        final_concepts.append(
            Concept(
                label=label,
                description=draft.description.strip(),
                topic=(draft.topic or "General").strip(),
                source_chunk_indices=[
                    draft.source_chunk_index
                ],
                locators=[locator] if locator else [],
            )
        )

    return final_concepts


# ============================================================
# QUESTION GENERATION
# ============================================================

@with_retry(
    retries=1,
    custom_exc=QuestionGenerationError,
)
async def generate_question(
    concept: Concept,
    difficulty: Difficulty,
    previous_questions: list[str] | None = None,
    question_type: str = "mcq",
) -> QuestionDraft:

    if client is None:

        raise QuestionGenerationError(
            "Gemini client is not initialized. "
            "Check GEMINI_API_KEY."
        )

    previous_questions = (
        previous_questions or []
    )

    # Keep the prompt from becoming unnecessarily huge.
    previous_questions = previous_questions[
        -20:
    ]

    if previous_questions:

        previous_text = "\n".join(
            f"- {question}"
            for question in previous_questions
        )

        uniqueness_instruction = f"""
IMPORTANT — QUESTIONS ALREADY ASKED

Do NOT repeat, paraphrase, or lightly modify any
of these previously asked questions:

{previous_text}

The new question MUST test a different aspect,
relationship, example, application, or reasoning
path within the concept.
"""

    else:

        uniqueness_instruction = """
There are no previously asked questions.
Generate an original question.
"""

    type_instructions = {
        "mcq": (
            "Create ONE multiple-choice question.\n\n"
            "Rules:\n"
            "1. Provide exactly 4 options.\n"
            "2. Exactly ONE option must have is_correct=true.\n"
            "3. The other THREE options must have is_correct=false.\n"
            "4. The correct answer must genuinely answer the question.\n"
            "5. Wrong options should represent realistic misconceptions.\n"
        ),
        "short_answer": (
            "Create ONE short-answer question (a one-sentence or "
            "one-phrase answer, not multiple choice).\n\n"
            "Rules:\n"
            "1. Leave 'options' as an empty list.\n"
            "2. Put the model answer in 'expected_answer'.\n"
            "3. The answer should be gradeable by meaning, not exact "
            "wording.\n"
        ),
        "numerical": (
            "Create ONE numerical problem with a single numeric answer.\n\n"
            "Rules:\n"
            "1. Leave 'options' as an empty list.\n"
            "2. Put the correct numeric value (as a string) in "
            "'expected_answer'.\n"
            "3. Put a reasonable acceptable tolerance in "
            "'numerical_tolerance' (e.g. 0.01 for an exact integer "
            "answer, larger for an approximate one).\n"
        ),
    }

    prompt = f"""
{type_instructions.get(question_type, type_instructions["mcq"])}

Concept:
{concept.label}

Concept description:
{concept.description}

Required difficulty:
{difficulty}

{uniqueness_instruction}

General rules:

1. Use clear educational language.
2. Match the requested difficulty.
3. Do not copy or paraphrase previous questions.
4. Test a different aspect of the concept when possible.
5. Set question_type to "{question_type}".
6. Return only the structured JSON response.
"""

    response = await client.aio.models.generate_content(

        model=MODEL_TEXT,

        contents=prompt,

        config=genai.types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=QuestionDraft,
            temperature=0.65,
        ),
    )

    question = QuestionDraft.model_validate_json(
        response.text
    )

    question.topic = concept.topic
    question.concept_label = concept.label
    question.source_locator = (
        concept.locators[0] if concept.locators else ""
    )

    # --------------------------------------------------------
    # Defensive validation
    # --------------------------------------------------------

    if question_type == "mcq":

        if len(question.options) != 4:

            raise QuestionGenerationError(
                "Gemini returned a question with "
                f"{len(question.options)} options instead of 4."
            )

        correct_count = sum(
            option.is_correct
            for option in question.options
        )

        if correct_count != 1:

            raise QuestionGenerationError(
                "Gemini returned a question without "
                "exactly one correct answer."
            )

    else:

        if not question.expected_answer:

            raise QuestionGenerationError(
                f"Gemini returned a {question_type} question "
                "without an expected_answer."
            )

    # Make sure Gemini actually respected the
    # requested difficulty.
    if question.difficulty != difficulty:

        raise QuestionGenerationError(
            "Gemini returned difficulty "
            f"'{question.difficulty}' instead of "
            f"requested '{difficulty}'."
        )

    # --------------------------------------------------------
    # Backend duplicate guard
    # --------------------------------------------------------

    normalized_new_question = (
        " ".join(
            question.question_text
            .lower()
            .split()
        )
    )

    for old_question in previous_questions:

        normalized_old_question = (
            " ".join(
                old_question
                .lower()
                .split()
            )
        )

        if (
            normalized_new_question
            == normalized_old_question
        ):

            raise QuestionGenerationError(
                "Gemini generated a duplicate question."
            )

    return question


# ============================================================
# QUESTION VALIDATION
# ============================================================

@with_retry(
    retries=1,
    custom_exc=QuestionGenerationError,
)
async def validate_question(
    question: QuestionDraft,
) -> ValidationResult:

    if client is None:

        raise QuestionGenerationError(
            "Gemini client is not initialized."
        )

    if question.question_type != "mcq":

        prompt = (
            "You are an expert educational question validator.\n\n"

            "Check whether this question has exactly one "
            "unambiguous correct answer, and that the given "
            "expected answer is actually correct.\n\n"

            f"Question type: {question.question_type}\n"
            f"Question:\n{question.question_text}\n\n"
            f"Expected answer:\n{question.expected_answer}\n\n"

            "Return passed=true only if the expected answer is "
            "clearly and unambiguously correct for this question."
        )

    else:

        correct_indices = [
            i
            for i, option in enumerate(
                question.options
            )
            if option.is_correct
        ]

        prompt = (
            "You are an expert educational question validator.\n\n"

            "Check whether this multiple-choice question "
            "has exactly one unambiguously correct answer.\n\n"

            f"Question:\n{question.question_text}\n\n"

            "Options:\n"
            f"{[option.text for option in question.options]}\n\n"

            "Intended correct answer index:\n"
            f"{correct_indices}\n\n"

            "Return passed=true only if the intended answer "
            "is clearly correct and all distractors are clearly wrong."
        )

    response = await client.aio.models.generate_content(

        model=MODEL_TEXT,

        contents=prompt,

        config=genai.types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=ValidationResult,
            temperature=0.1,
        ),
    )

    return ValidationResult.model_validate_json(
        response.text
    )


# ============================================================
# GENERATE + VALIDATE
# ============================================================

async def generate_validated_question(
    concept: Concept,
    difficulty: Difficulty,
    previous_questions: list[str] | None = None,
    question_type: str = "mcq",
) -> QuestionDraft:

    previous_questions = (
        previous_questions or []
    )

    last_error = None

    # Multiple attempts are important because Gemini
    # may occasionally produce a duplicate or invalid question.
    for attempt in range(4):

        try:

            question = await generate_question(
                concept=concept,
                difficulty=difficulty,
                previous_questions=previous_questions,
                question_type=question_type,
            )

            validation = await validate_question(
                question
            )

            if validation.passed:

                # Randomize answer positions after validation.
                # The is_correct flag stays attached to each option.

                random.shuffle(
                    question.options
                )

                return question

            last_error = (
                "Question failed validation: "
                f"{validation.notes}"
            )

        except Exception as e:

            last_error = str(e)

    raise ValidationFailedError(
        f"Could not generate a valid unique question "
        f"for '{concept.label}' at {difficulty} difficulty. "
        f"{last_error}"
    )


# ============================================================
# WRONG-ANSWER EXPLANATION
# ============================================================

@with_retry(retries=2)
async def generate_explanation(
    question: QuestionDraft,
    selected_option_index: int,
) -> str:

    selected_option = question.options[
        selected_option_index
    ]

    if selected_option.is_correct:

        return (
            "That is the correct answer! Well done."
        )

    misconception = (
        selected_option.misconception_tag
        or "general misunderstanding"
    )

    prompt = (
        "A student answered this question incorrectly.\n\n"
        f"Question: {question.question_text}\n"
        f"They selected: {selected_option.text}\n"
        f"Misconception: {misconception}\n\n"
        "Write a short, encouraging explanation "
        "of why the correct answer is right."
    )

    response = await client.aio.models.generate_content(
        model=MODEL_TEXT,
        contents=prompt,
        config=genai.types.GenerateContentConfig(
            temperature=0.3
        ),
    )

    return response.text.strip()


# ============================================================
# END-OF-QUIZ SUMMARY
# ============================================================

@with_retry(retries=2)
async def generate_quiz_summary(
    report_data: QuizReportData,
) -> str:

    data_str = "Concepts Tested:\n"

    for concept in report_data.concepts:

        data_str += (
            f"- {concept.concept_label}: "
            f"Ended at {concept.final_difficulty.upper()} "
            f"difficulty. "
            f"Accuracy: {concept.accuracy:.0%}.\n"
        )

    data_str += (
        "\nWeakest Concepts: "
        + ", ".join(
            report_data.weakest_concept_labels
        )
    )

    prompt = (
        "You are an empathetic professional tutor.\n"
        "Based on the following quiz report, write "
        "one encouraging paragraph summarizing "
        "the student's performance.\n\n"
        f"{data_str}"
    )

    response = await client.aio.models.generate_content(
        model=MODEL_TEXT,
        contents=prompt,
        config=genai.types.GenerateContentConfig(
            temperature=0.5
        ),
    )

    return response.text.strip()


# ============================================================
# STRETCH: PAST PAPER CALIBRATION
# ============================================================

@with_retry(retries=1)
async def analyze_past_paper(
    pdf_bytes: bytes,
) -> CalibrationEstimate:

    chunks = extract_and_chunk(
        pdf_bytes
    )[:3]

    combined_text = "\n\n".join(
        chunk.text
        for chunk in chunks
    )

    prompt = (
        "Review this exam excerpt and estimate "
        "its difficulty as easy, medium, or hard.\n\n"
        f"{combined_text}"
    )

    response = await client.aio.models.generate_content(
        model=MODEL_TEXT,
        contents=prompt,
        config=genai.types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=CalibrationEstimate,
            temperature=0.2,
        ),
    )

    return CalibrationEstimate.model_validate_json(
        response.text
    )


# ============================================================
# EMBEDDINGS / RETRIEVAL
# ============================================================

EMBEDDING_MODEL = "gemini-embedding-001"

GROUNDING_THRESHOLD = 0.55


async def embed_texts(
    texts: list[str],
) -> list[list[float]]:

    if client is None:
        raise ConceptExtractionError(
            "Gemini client is not initialized. Check GEMINI_API_KEY."
        )

    if not texts:
        return []

    response = await client.aio.models.embed_content(
        model=EMBEDDING_MODEL,
        contents=texts,
    )

    return [
        list(embedding.values)
        for embedding in response.embeddings
    ]


def _cosine_similarity(
    a: list[float],
    b: list[float],
) -> float:

    dot = sum(x * y for x, y in zip(a, b))

    norm_a = sum(x * x for x in a) ** 0.5
    norm_b = sum(y * y for y in b) ** 0.5

    if norm_a == 0 or norm_b == 0:
        return 0.0

    return dot / (norm_a * norm_b)


async def retrieve_relevant_chunks(
    chunks: list[Chunk],
    chunk_embeddings: list[list[float]],
    query: str,
    k: int = 4,
) -> list[tuple[Chunk, float]]:

    if not chunks or not chunk_embeddings:
        return []

    query_embedding = (
        await embed_texts([query])
    )[0]

    scored = [
        (chunk, _cosine_similarity(query_embedding, embedding))
        for chunk, embedding in zip(chunks, chunk_embeddings)
    ]

    scored.sort(key=lambda pair: pair[1], reverse=True)

    return scored[:k]


# ============================================================
# GROUNDED TUTOR CHAT (source-cited, with refusal)
# ============================================================

@with_retry(retries=2)
async def answer_grounded_question(
    chunks: list[Chunk],
    chunk_embeddings: list[list[float]],
    conversation_history: list[ChatTurn],
    student_message: str,
) -> GroundedChatResponse:

    if client is None:
        raise QuestionGenerationError(
            "Gemini client is not initialized."
        )

    top_matches = await retrieve_relevant_chunks(
        chunks,
        chunk_embeddings,
        student_message,
        k=4,
    )

    if not top_matches or top_matches[0][1] < GROUNDING_THRESHOLD:

        return GroundedChatResponse(
            answer=(
                "I can't find anything in your uploaded material "
                "that covers this. Try rephrasing, or ask about a "
                "topic that's in your notes."
            ),
            is_grounded=False,
            citations=[],
        )

    context_blocks = "\n\n".join(
        f"[Source: {chunk.locator}]\n{chunk.text}"
        for chunk, _score in top_matches
    )

    history_text = "\n".join(
        f"{turn.role}: {turn.text}"
        for turn in conversation_history[-6:]
    )

    prompt = (
        "You are a study tutor that must answer ONLY using the "
        "source excerpts below. Cite the exact [Source: ...] tag(s) "
        "you used inline in your answer. If the excerpts don't "
        "actually answer the question, say so instead of guessing.\n\n"
        f"Source excerpts:\n{context_blocks}\n\n"
        f"Conversation so far:\n{history_text}\n\n"
        f"Student question:\n{student_message}"
    )

    response = await client.aio.models.generate_content(
        model=MODEL_TEXT,
        contents=prompt,
        config=genai.types.GenerateContentConfig(
            temperature=0.3,
        ),
    )

    answer_text = (response.text or "").strip()

    citations = [
        Citation(
            chunk_index=chunk.chunk_index,
            locator=chunk.locator,
            excerpt=chunk.text[:280],
        )
        for chunk, _score in top_matches
    ]

    return GroundedChatResponse(
        answer=answer_text,
        is_grounded=True,
        citations=citations,
    )


# ============================================================
# ANSWER VERIFICATION (short-answer / numerical / cross-check)
# ============================================================

@with_retry(retries=1)
async def verify_answer(
    question: QuestionDraft,
    student_answer: str,
) -> AnswerVerification:

    if question.question_type == "numerical":

        try:

            given = float(student_answer)
            expected = float(question.expected_answer or "nan")
            tolerance = question.numerical_tolerance or 0.01

            is_correct = abs(given - expected) <= tolerance

            return AnswerVerification(
                is_correct=is_correct,
                confidence=1.0,
                feedback=(
                    "Correct."
                    if is_correct
                    else f"Expected {question.expected_answer}."
                ),
                source_locator=question.source_locator,
            )

        except (ValueError, TypeError):
            pass

    if client is None:
        raise QuestionGenerationError(
            "Gemini client is not initialized."
        )

    prompt = (
        "You are an independent answer grader. Judge ONLY whether "
        "the student's answer is substantively correct, even if "
        "worded differently from the expected answer.\n\n"
        f"Question:\n{question.question_text}\n\n"
        f"Expected answer:\n{question.expected_answer}\n\n"
        f"Student answer:\n{student_answer}\n\n"
        "Give brief, specific feedback."
    )

    response = await client.aio.models.generate_content(
        model=MODEL_TEXT,
        contents=prompt,
        config=genai.types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=AnswerVerification,
            temperature=0.1,
        ),
    )

    result = AnswerVerification.model_validate_json(
        response.text
    )

    result.source_locator = question.source_locator

    return result


# ============================================================
# TUTOR FOLLOW-UP CHAT
# ============================================================

@with_retry(retries=2)
async def tutor_followup(
    concept: Concept,
    conversation_history: list[ChatTurn],
    student_message: str,
) -> str:

    contents = [
        genai.types.Content(
            role=(
                "user"
                if turn.role == "user"
                else "model"
            ),
            parts=[
                genai.types.Part.from_text(
                    text=turn.text
                )
            ],
        )
        for turn in conversation_history
    ]

    system_instruction = (
        f"You are a helpful tutor explaining "
        f"'{concept.label}' "
        f"({concept.description}). "
        "Keep answers concise, encouraging, "
        "and in plain language."
    )

    contents.append(
        genai.types.Content(
            role="user",
            parts=[
                genai.types.Part.from_text(
                    text=(
                        f"{system_instruction}\n\n"
                        f"Student: {student_message}"
                    )
                )
            ],
        )
    )

    response = await client.aio.models.generate_content(
        model=MODEL_TEXT,
        contents=contents,
        config=genai.types.GenerateContentConfig(
            temperature=0.6
        ),
    )

    return response.text.strip()